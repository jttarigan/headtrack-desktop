"""Excel workbooks for the side paper's data figures (LAB.md, Paper finishing): the plotted values plus a native Excel
chart, so a figure can be remade in Excel.

    python3 headtrack/paper/make_xlsx.py          (from the repository root; system Python, which has openpyxl)

Reads headtrack/paper/results/fig_data.json (written by headtrack/analysis/paper_results.py) and writes
headtrack/paper/results/xlsx/figure3_turn.xlsx ... figure6_drift.xlsx, numbered as in the manuscript. Figures 1
(architecture) and 2 (photos) show no data. Typed-in values are blue; medians are Excel formulas (black).
"""
import json
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RES = os.path.join(ROOT, "headtrack", "paper", "results")
OUT = os.path.join(RES, "xlsx")

# Same colours as paper_results.py, by label.
COLOR = {"Reference (iPhone)": "222222", "AirPods Max 2": "1F77B4", "AirPods Pro (1st gen.)": "9467BD",
         "MediaPipe": "FF7F0E", "Apple Vision": "2CA02C"}
FONT, BOLD, INPUT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), \
    Font(name="Arial", size=10, color="0000FF")


def new_book(title, notes):
    """Workbook with a README sheet (title + notes) and an empty Data sheet."""
    wb = Workbook()
    rd = wb.active
    rd.title = "README"
    rd["A1"] = title
    rd["A1"].font = Font(name="Arial", size=12, bold=True)
    for i, line in enumerate(notes, start=3):
        rd.cell(i, 1, line).font = FONT
    rd.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Data")
    wb.calculation.fullCalcOnLoad = True
    return wb, ws


def header(ws, row, names, width=16):
    for c, n in enumerate(names, start=1):
        ws.cell(row, c, n).font = BOLD
        ws.column_dimensions[get_column_letter(c)].width = width


def put(ws, row, col, value, fmt=None):
    """Text and formulas in black, typed-in numbers in blue."""
    cell = ws.cell(row, col, value)
    cell.font = FONT if isinstance(value, str) else INPUT
    if fmt:
        cell.number_format = fmt
    return cell


def colour(series, label, line=True, marker=None):
    c = COLOR.get(label, "555555")
    if line:
        series.graphicalProperties.line.solidFill = c
        series.graphicalProperties.line.width = 19050
    else:
        series.graphicalProperties.line.noFill = True
    if marker:
        series.marker.symbol = marker
        series.marker.size = 5
        series.marker.graphicalProperties.solidFill = c
        series.marker.graphicalProperties.line.solidFill = c
    else:
        series.marker.symbol = "none"
    series.smooth = False


def figure3_turn(fd):
    wb, ws = new_book("Figure 3: one quick turn (session %s)" % fd["session"], [
        "Yaw of each tracker during one quick turn to the left and back, as in Figure 3 of the manuscript.",
        "Reference: iPhone ARKit yaw at capture time. Trackers: yaw as the application received it (latest sample "
        "held until the next one arrives).",
        "Each curve is shifted so that it matches the reference during a held pose, as in the paper's figure.",
        "Columns per curve: time (ms) and yaw (deg). The 'step' columns repeat each tracker sample until the next "
        "arrives; the chart draws those, so the curves are steps as in the paper.",
        "Source: headtrack/paper/results/fig_data.json, written by headtrack/analysis/paper_results.py.",
    ])
    ch = ScatterChart()
    ch.title, ch.style = "One quick turn", 13
    ch.x_axis.title, ch.y_axis.title = "time (ms)", "yaw (deg)"
    ch.width, ch.height = 22, 10
    col = 1
    for s in fd["series"]:
        n = len(s["t_ms"])
        ws.cell(1, col, s["label"] + " time (ms)").font = BOLD
        ws.cell(1, col + 1, s["label"] + " yaw (deg)").font = BOLD
        for i, (t, y) in enumerate(zip(s["t_ms"], s["yaw_deg"]), start=2):
            put(ws, i, col, t, "0.0")
            put(ws, i, col + 1, y, "0.00")
        xs, ys, last = col, col + 1, n + 1
        if s["steps"]:  # step-expanded copy for the chart: (t_i, y_i), (t_i+1, y_i), ...
            ws.cell(1, col + 2, s["label"] + " step time (ms)").font = BOLD
            ws.cell(1, col + 3, s["label"] + " step yaw (deg)").font = BOLD
            r = 2
            for i in range(n):
                for it in ((i, i + 1) if i + 1 < n else (i,)):
                    put(ws, r, col + 2, "=%s%d" % (get_column_letter(col), 2 + it), "0.0")
                    put(ws, r, col + 3, "=%s%d" % (get_column_letter(col + 1), 2 + i), "0.00")
                    r += 1
            xs, ys, last = col + 2, col + 3, r - 1
        se = Series(Reference(ws, min_col=ys, min_row=2, max_row=last), Reference(ws, min_col=xs, min_row=2, max_row=last),
                    title=s["label"] + (" (capture time)" if not s["steps"] else " (as received)"))
        colour(se, s["label"])
        ch.series.append(se)
        for c in range(col, col + (4 if s["steps"] else 2)):
            ws.column_dimensions[get_column_letter(c)].width = 14
        col += 5 if s["steps"] else 3
    cs = wb.create_sheet("Chart")
    cs.add_chart(ch, "A1")
    return wb


def figure4_latency(fd):
    wb, ws = new_book("Figure 4: motion-to-data delay per tracker", [
        "Bars: median motion-to-data delay per tracker over the latency sessions (Excel MEDIAN of the session values).",
        "Markers in the paper's figure: the per-session medians below, one shape per recording day.",
        "Grey band in the paper's figure: %d to %d ms, thresholds of the most sensitive listeners." % tuple(fd["band_ms"]),
        "Source: headtrack/paper/results/fig_data.json, written by headtrack/analysis/paper_results.py.",
    ])
    tr = fd["trackers"]
    sessions = [(x["label"], x["day"]) for x in tr[0]["sessions"]]
    header(ws, 1, ["session", "day"] + [t["label"] + " (ms)" for t in tr])
    for i, (lab, day) in enumerate(sessions, start=2):
        put(ws, i, 1, lab)
        put(ws, i, 2, day, "0")
        for k, t in enumerate(tr):
            v = {x["label"]: x["ms"] for x in t["sessions"]}.get(lab)
            if v is not None:
                put(ws, i, 3 + k, v, "0.0")
    last = len(sessions) + 1
    top = last + 3
    header(ws, top, ["tracker", "median (ms)", "band from (ms)", "band to (ms)"])
    for k, t in enumerate(tr):
        r = top + 1 + k
        c = get_column_letter(3 + k)
        put(ws, r, 1, t["label"])
        put(ws, r, 2, "=MEDIAN(%s2:%s%d)" % (c, c, last), "0.0")
        put(ws, r, 3, fd["band_ms"][0], "0")
        put(ws, r, 4, fd["band_ms"][1], "0")
    ch = BarChart()
    ch.type, ch.style = "bar", 10
    ch.title = "Motion-to-data delay, median"
    ch.x_axis.title, ch.y_axis.title = "tracker", "delay (ms)"
    ch.add_data(Reference(ws, min_col=2, min_row=top, max_row=top + len(tr)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=top + len(tr)))
    ch.x_axis.scaling.orientation = "maxMin"  # first tracker on top, as in the paper
    ch.legend = None
    for k, t in enumerate(tr):
        pt = DataPoint(idx=k)
        pt.graphicalProperties.solidFill = COLOR.get(t["label"], "555555")
        ch.series[0].dPt.append(pt)
    ch.width, ch.height = 18, 8
    # second chart: per-session values by tracker (the paper's markers)
    ch2 = BarChart()
    ch2.type, ch2.style = "bar", 10
    ch2.title = "Per-session medians (sessions of day 1, then day 2)"
    ch2.y_axis.title = "delay (ms)"
    ch2.add_data(Reference(ws, min_col=3, max_col=2 + len(tr), min_row=1, max_row=last), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ch2.x_axis.scaling.orientation = "maxMin"
    for k, t in enumerate(tr):
        ch2.series[k].graphicalProperties.solidFill = COLOR.get(t["label"], "555555")
    ch2.width, ch2.height = 18, 12
    cs = wb.create_sheet("Chart")
    cs.add_chart(ch, "A1")
    cs.add_chart(ch2, "A18")
    return wb


def figure5_range(fd):
    wb, ws = new_book("Figure 5: tracking range by reference yaw", [
        "(a) Share of time each webcam tracker reported a pose, per bin of the reference's absolute yaw (%).",
        "(b) Median absolute yaw error against the reference per bin (deg); empty where a bin had fewer than 20 samples.",
        "Source: headtrack/paper/results/fig_data.json, written by headtrack/analysis/paper_results.py.",
    ])
    bins = [b.replace("-", "–") for b in fd["bins"]]
    sh, er = fd["share_pct"], fd["yaw_error_deg"]
    header(ws, 1, ["|yaw| bin (deg)"] + ["%s pose (%%)" % n for n in sh] + ["", "|yaw| bin (deg)"] +
           ["%s |error| (deg)" % n for n in er], width=18)
    e0 = len(sh) + 3
    for i, b in enumerate(bins, start=2):
        put(ws, i, 1, b)
        put(ws, i, e0, b)
        for k, n in enumerate(sh):
            put(ws, i, 2 + k, sh[n][i - 2], "0.0")
        for k, n in enumerate(er):
            if er[n][i - 2] is not None:
                put(ws, i, e0 + 1 + k, er[n][i - 2], "0.00")
    last = len(bins) + 1
    a = BarChart()
    a.type, a.style, a.grouping = "col", 10, "clustered"
    a.title, a.y_axis.title, a.x_axis.title = "(a) time with a pose", "time with a pose (%)", "|yaw| from the reference (deg)"
    a.add_data(Reference(ws, min_col=2, max_col=1 + len(sh), min_row=1, max_row=last), titles_from_data=True)
    a.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    a.y_axis.scaling.min, a.y_axis.scaling.max = 0, 100
    for k, n in enumerate(sh):
        a.series[k].graphicalProperties.solidFill = COLOR.get(n, "555555")
    b = LineChart()
    b.style = 12
    b.title, b.y_axis.title, b.x_axis.title = "(b) median |yaw error|", "median |yaw error| (deg)", "|yaw| from the reference (deg)"
    b.add_data(Reference(ws, min_col=e0 + 1, max_col=e0 + len(er), min_row=1, max_row=last), titles_from_data=True)
    b.set_categories(Reference(ws, min_col=e0, min_row=2, max_row=last))
    for k, n in enumerate(er):
        colour(b.series[k], n, marker="circle")
    for c in (a, b):
        c.width, c.height = 16, 9
    cs = wb.create_sheet("Chart")
    cs.add_chart(a, "A1")
    cs.add_chart(b, "K1")
    return wb


def figure6_drift(fd):
    notes = [
        "Headphone yaw minus reference yaw, per-minute median, relative to the first minute (deg), for each drift session;",
        "MediaPipe yaw minus reference yaw on the same minutes for comparison. Minute = midpoint of each one-minute window.",
    ]
    for d in fd:
        for j in d["heading_jump"]:
            notes.append("Session %s: heading jump of %.0f deg at minute %.2f; its curve stops before it (dotted line in "
                         "the paper's figure)." % (d["session"], j["deg"], j["min"]))
    notes.append("Source: headtrack/paper/results/fig_data.json, written by headtrack/analysis/paper_results.py.")
    wb, ws = new_book("Figure 6: headphone yaw drift", notes)
    ch = ScatterChart()
    ch.title, ch.style = "Yaw drift during desk work", 13
    ch.x_axis.title, ch.y_axis.title = "time (min)", "yaw difference (deg)"
    ch.width, ch.height = 22, 11
    marks = ("circle", "triangle", "diamond", "square")
    col = 1
    for k, d in enumerate(fd):
        hp = "%s minus reference (%s)" % (d["headphones"], d["session"])
        mp = "MediaPipe minus reference (%s)" % d["session"]
        for c, n in enumerate(("minute (%s)" % d["session"], hp, mp)):
            ws.cell(1, col + c, n).font = BOLD
            ws.column_dimensions[get_column_letter(col + c)].width = 22
        for i, x in enumerate(d["minute_mid"], start=2):
            put(ws, i, col, x, "0.0")
            if d["headphones_minus_ref_deg"][i - 2] is not None:
                put(ws, i, col + 1, d["headphones_minus_ref_deg"][i - 2], "0.00")
            if d["mediapipe_minus_ref_deg"] and d["mediapipe_minus_ref_deg"][i - 2] is not None:
                put(ws, i, col + 2, d["mediapipe_minus_ref_deg"][i - 2], "0.00")
        last = len(d["minute_mid"]) + 1
        for c, title, lab in ((col + 1, hp, d["headphones"]), (col + 2, mp, "MediaPipe")):
            if c == col + 2 and not d["mediapipe_minus_ref_deg"]:
                continue
            se = Series(Reference(ws, min_col=c, min_row=2, max_row=last), Reference(ws, min_col=col, min_row=2, max_row=last),
                        title=title)
            colour(se, lab, marker=marks[k % 4])
            if k:
                se.graphicalProperties.line.dashStyle = "dash"
            ch.series.append(se)
        col += 4
    cs = wb.create_sheet("Chart")
    cs.add_chart(ch, "A1")
    return wb


def main():
    fd = json.load(open(os.path.join(RES, "fig_data.json")))
    os.makedirs(OUT, exist_ok=True)
    books = {"figure3_turn": (figure3_turn, "turn"), "figure4_latency": (figure4_latency, "latency"),
             "figure5_range": (figure5_range, "range"), "figure6_drift": (figure6_drift, "drift")}
    for name, (make, key) in books.items():
        if fd.get(key):
            path = os.path.join(OUT, name + ".xlsx")
            make(fd[key]).save(path)
            print("wrote", os.path.relpath(path, ROOT))


if __name__ == "__main__":
    main()
